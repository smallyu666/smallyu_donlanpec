"""单独法兰报告、图纸生成流程。

此模块不依赖产品 ID，不读取产品设计活动表。参数总览、报告和图纸
直接从程序根目录的 ``法兰独立计算Output.json`` 读取数据。
"""
from __future__ import annotations

import ctypes

try:
    import winreg
except ImportError:
    winreg = None

from PyQt5.QtCore import QStandardPaths

import json
import os
import re
import shutil
import sys
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from PyQt5.QtWidgets import QFileDialog, QMessageBox


FLANGE_MODULE_NAME = "法兰"
JSON_FILENAME = "flange_calculate.json"
FLANGE_OUTPUT_JSON_FILENAME = "法兰独立计算Output.json"
FLANGE_OUTPUT_MODULE_NAMES = ("法兰独立计算", FLANGE_MODULE_NAME)
INTERMEDIATE_TEMPLATE_FILENAME = "强度计算元件输出参数表_法兰.xlsx"
REPORT_TEMPLATE_FILENAME = "计算报告（法兰）.xlsx"
DESKTOP_REPORT_PATTERN = re.compile(
    r"^法兰计算报告(?P<timestamp>\d{4}-\d{2}-\d{2}-\d{2}-\d{2}-\d{2})\.(?:xlsx|xlsm)$",
    re.IGNORECASE,
)
GENERIC_REPORT_LABELS = {"", "值", "允许值", "许用值", "校核结果", "结论", "计算结果", "待补充参数"}


def _program_root(program_root: Optional[str] = None) -> Path:
    if program_root:
        return Path(program_root).resolve()
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parents[2]


def _normalise_key(value: Any) -> str:
    text = str(value or "").strip()
    return re.sub(r"\s+", "", text.replace("（", "(").replace("）", ")"))


def _normalise_header(value: Any) -> str:
    """表头识别不区分大小写；计算参数名必须保留 F/f 等英文字母大小写。"""
    return _normalise_key(value).casefold()


def _stringify_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value).strip()


def _safe_path(value: Any) -> Optional[Path]:
    """展开环境变量和用户目录，并转换为 Path。"""
    if value is None:
        return None

    text = str(value).strip().strip('"')
    if not text:
        return None

    try:
        text = os.path.expandvars(text)
        text = os.path.expanduser(text)
        return Path(text)
    except Exception:
        return None


def _qt_desktop_directory() -> Optional[Path]:
    """
    使用 Qt 获取系统桌面路径。

    Windows 下通常能够识别：
    - 普通桌面
    - OneDrive 重定向桌面
    - 企业组策略重定向桌面
    """
    try:
        desktop = QStandardPaths.writableLocation(
            QStandardPaths.DesktopLocation
        )
        return _safe_path(desktop)
    except Exception as error:
        print(f"[桌面路径] Qt 获取失败：{error}")
        return None


def _windows_shell_desktop_directory() -> Optional[Path]:
    """
    使用 Windows Shell 获取当前用户桌面目录。

    CSIDL_DESKTOPDIRECTORY = 0x0010
    """
    if os.name != "nt":
        return None

    try:
        buffer = ctypes.create_unicode_buffer(32768)

        result = ctypes.windll.shell32.SHGetFolderPathW(
            None,
            0x0010,  # CSIDL_DESKTOPDIRECTORY
            None,
            0,
            buffer,
        )

        if result == 0 and buffer.value:
            return _safe_path(buffer.value)

    except Exception as error:
        print(f"[桌面路径] Windows Shell 获取失败：{error}")

    return None


def _registry_desktop_directory() -> Optional[Path]:
    """
    从注册表读取当前用户桌面路径。

    该位置一般可以识别 OneDrive、网络重定向桌面等情况。
    """
    if os.name != "nt" or winreg is None:
        return None

    registry_path = (
        r"Software\Microsoft\Windows"
        r"\CurrentVersion\Explorer\User Shell Folders"
    )

    try:
        with winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            registry_path,
            0,
            winreg.KEY_READ,
        ) as key:
            value, value_type = winreg.QueryValueEx(
                key,
                "Desktop",
            )

        if value_type == winreg.REG_EXPAND_SZ:
            value = os.path.expandvars(value)

        return _safe_path(value)

    except FileNotFoundError:
        print("[桌面路径] 注册表中没有 Desktop 项")
    except Exception as error:
        print(f"[桌面路径] 注册表读取失败：{error}")

    return None


def _desktop_directories() -> List[Path]:
    """
    获取可能的桌面目录。

    获取顺序：
    1. Qt 标准桌面路径
    2. Windows Shell 桌面路径
    3. Windows 注册表桌面路径
    4. USERPROFILE
    5. OneDrive 本地/个人/企业桌面
    6. 公共桌面
    """
    candidates: List[Optional[Path]] = []

    # 1. 系统接口优先
    candidates.append(_qt_desktop_directory())
    candidates.append(_windows_shell_desktop_directory())
    candidates.append(_registry_desktop_directory())

    # 2. 普通用户目录兜底
    home = Path.home()
    candidates.extend(
        [
            home / "Desktop",
            home / "桌面",
        ]
    )

    user_profile = os.environ.get("USERPROFILE")
    if user_profile:
        profile_path = _safe_path(user_profile)
        if profile_path:
            candidates.extend(
                [
                    profile_path / "Desktop",
                    profile_path / "桌面",
                ]
            )

    # 3. OneDrive 目录兜底
    for environment_name in (
        "OneDrive",
        "OneDriveConsumer",
        "OneDriveCommercial",
    ):
        one_drive_value = os.environ.get(environment_name)
        if not one_drive_value:
            continue

        one_drive_path = _safe_path(one_drive_value)
        if one_drive_path:
            candidates.extend(
                [
                    one_drive_path / "Desktop",
                    one_drive_path / "桌面",
                ]
            )

    # 4. 公共桌面兜底
    public_path = _safe_path(os.environ.get("PUBLIC"))
    if public_path:
        candidates.extend(
            [
                public_path / "Desktop",
                public_path / "桌面",
            ]
        )

    desktop_directories: List[Path] = []
    seen = set()

    for path in candidates:
        if path is None:
            continue

        try:
            path_key = os.path.normcase(
                os.path.normpath(str(path))
            )

            if path_key in seen:
                continue

            seen.add(path_key)

            if path.is_dir():
                desktop_directories.append(path)
                print(f"[桌面路径] 有效目录：{path}")
            else:
                print(f"[桌面路径] 目录不存在：{path}")

        except OSError as error:
            print(f"[桌面路径] 无法访问 {path}：{error}")

    return desktop_directories

def timestamp_from_report_path(report_path: Path) -> str:
    """优先使用计算程序写在文件名中的时间；无法解析时回退到文件修改时间。"""
    match = DESKTOP_REPORT_PATTERN.fullmatch(report_path.name)
    if match:
        return match.group("timestamp")
    return datetime.fromtimestamp(report_path.stat().st_mtime).strftime("%Y-%m-%d-%H-%M-%S")


def find_desktop_flange_report(parent=None) -> Path:
    """
    查找桌面最新的法兰计算报告。

    自动查找失败时，允许用户手动选择 Excel，
    避免因为特殊桌面重定向导致整个流程无法使用。
    """
    desktop_directories = _desktop_directories()
    matches: List[Path] = []

    for desktop in desktop_directories:
        try:
            print(f"[法兰报告] 正在检查：{desktop}")

            for item in desktop.iterdir():
                try:
                    if (
                        item.is_file()
                        and DESKTOP_REPORT_PATTERN.fullmatch(item.name)
                    ):
                        matches.append(item)
                        print(f"[法兰报告] 找到：{item}")
                except OSError:
                    continue

        except OSError as error:
            print(
                f"[法兰报告] 无法读取桌面目录 "
                f"{desktop}：{error}"
            )

    if matches:
        valid_matches = []

        for item in matches:
            try:
                valid_matches.append(item)
            except OSError:
                continue

        if valid_matches:
            latest_report = max(
                valid_matches,
                key=lambda item: (
                    timestamp_from_report_path(item),
                    item.stat().st_mtime,
                ),
            )

            print(f"[法兰报告] 最终使用：{latest_report}")
            return latest_report

    # 自动搜索失败，允许用户手动选择
    if desktop_directories:
        initial_directory = str(desktop_directories[0])
    else:
        initial_directory = str(Path.home())

    selected_file, _ = QFileDialog.getOpenFileName(
        parent,
        "未自动找到法兰报告，请手动选择",
        initial_directory,
        "法兰计算报告 (*.xlsx *.xlsm);;"
        "Excel 文件 (*.xlsx *.xlsm);;"
        "所有文件 (*.*)",
    )

    if selected_file:
        selected_path = Path(selected_file)

        if not selected_path.is_file():
            raise FileNotFoundError(
                f"选择的文件不存在：\n{selected_path}"
            )

        print(f"[法兰报告] 用户手动选择：{selected_path}")
        return selected_path

    searched_text = "\n".join(
        str(path) for path in desktop_directories
    )

    if not searched_text:
        searched_text = "未能获取任何有效桌面路径"

    raise FileNotFoundError(
        "未找到法兰计算报告。\n\n"
        "程序已经检查以下目录：\n"
        f"{searched_text}\n\n"
        "需要的默认文件名格式为：\n"
        "法兰计算报告YYYY-MM-DD-HH-MM-SS.xlsx"
    )

def _find_parameter_columns(sheet) -> Tuple[int, int, int]:
    """识别 Excel 的参数名、参数值列，支持表头不在第一行。"""
    for row_index in range(1, min(sheet.max_row, 12) + 1):
        headers = {
            _normalise_header(sheet.cell(row=row_index, column=column).value): column
            for column in range(1, sheet.max_column + 1)
        }
        name_column = next(
            (column for header, column in headers.items() if header in {"参数名", "参数名称", "名称"}),
            None,
        )
        value_column = next(
            (column for header, column in headers.items() if header in {"参数值", "值", "value"}),
            None,
        )
        if name_column and value_column:
            return row_index, name_column, value_column
    raise ValueError("未找到“参数名”和“参数值”两列，请检查法兰计算报告 Excel。")


def convert_flange_excel_to_json(excel_path: Path, json_path: Path) -> Dict[str, Any]:
    """将单独法兰 Excel 转为通用计算输出的 DictOutDatas JSON 格式。"""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = None
        header_row = name_column = value_column = None
        for candidate in workbook.worksheets:
            try:
                header_row, name_column, value_column = _find_parameter_columns(candidate)
                sheet = candidate
                break
            except ValueError:
                continue
        if sheet is None or header_row is None or name_column is None or value_column is None:
            raise ValueError("未找到“参数名”和“参数值”两列，请检查法兰计算报告 Excel。")

        datas: List[Dict[str, Any]] = []
        for row_index in range(header_row + 1, sheet.max_row + 1):
            name = _stringify_value(sheet.cell(row=row_index, column=name_column).value)
            if not name:
                continue
            datas.append(
                {
                    "Id": f"工况1：FL{len(datas) + 1}",
                    "Name": name,
                    "Value": _stringify_value(sheet.cell(row=row_index, column=value_column).value),
                    "Desc": name,
                    "ReportOutput": None,
                }
            )
    finally:
        workbook.close()

    if not datas:
        raise ValueError("法兰计算报告中没有可转换的参数数据。")

    result = {
        "Logs": [],
        "DictOutDatas": {
            FLANGE_MODULE_NAME: {
                "IsSuccess": not any("不合格" in item["Value"] for item in datas),
                "Datas": datas,
            }
        },
    }
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=4), encoding="utf-8")
    return result


def load_flange_output_json(root: Path) -> Tuple[Path, Dict[str, Any]]:
    """读取根目录法兰独立计算输出 JSON，并做最小格式校验。"""
    output_path = root / FLANGE_OUTPUT_JSON_FILENAME
    if not output_path.is_file():
        raise FileNotFoundError(
            "未找到单独法兰计算输出文件：\n"
            f"{output_path}\n\n"
            "请先在“单独法兰计算”窗口点击“计算”，生成 Output.json 后再导出。"
        )

    try:
        json_data = json.loads(output_path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as error:
        raise ValueError(
            f"{FLANGE_OUTPUT_JSON_FILENAME} 不是有效 JSON：{error}"
        ) from error

    if not isinstance(json_data, dict):
        raise ValueError(f"{FLANGE_OUTPUT_JSON_FILENAME} 的根节点必须是 JSON 对象。")

    _get_flange_datas(json_data)
    return output_path, json_data


def _get_flange_module_data(json_data: Dict[str, Any]) -> Dict[str, Any]:
    dict_out = json_data.get("DictOutDatas") or {}
    if not isinstance(dict_out, dict):
        raise ValueError(f"{FLANGE_OUTPUT_JSON_FILENAME} 的 DictOutDatas 格式不正确。")

    for module_name in FLANGE_OUTPUT_MODULE_NAMES:
        module_data = dict_out.get(module_name)
        if isinstance(module_data, dict) and isinstance(module_data.get("Datas"), list):
            return module_data

    for module_data in dict_out.values():
        if isinstance(module_data, dict) and isinstance(module_data.get("Datas"), list):
            return module_data

    raise ValueError(
        f"{FLANGE_OUTPUT_JSON_FILENAME} 中未找到法兰计算结果 Datas。"
    )


def _get_flange_datas(json_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    module_data = _get_flange_module_data(json_data)
    datas = module_data.get("Datas") or []
    if not isinstance(datas, list):
        raise ValueError(f"{FLANGE_OUTPUT_JSON_FILENAME} 的法兰参数格式不正确。")
    return [item for item in datas if isinstance(item, dict)]


def _parameter_value_map(datas: Iterable[Dict[str, Any]]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for item in datas:
        name = _stringify_value(item.get("Name"))
        if name and name not in result:
            result[name] = _stringify_value(item.get("Value"))
    return result


def _find_value(value_map: Dict[str, str], candidates: Sequence[str]) -> Optional[str]:
    for candidate in candidates:
        if candidate in value_map:
            return value_map[candidate]
    normalised_map = {_normalise_key(key): value for key, value in value_map.items()}
    for candidate in candidates:
        value = normalised_map.get(_normalise_key(candidate))
        if value is not None:
            return value
    return None


def _xlwings_sheet(workbook, sheet_name: str):
    try:
        return workbook.sheets[sheet_name]
    except Exception:
        return workbook.sheets[0]


def _try_generate_with_xlwings(
    template_path: Path,
    output_path: Path,
    fill_workbook,
    label: str,
) -> bool:
    """优先用 Excel 写入，保留模板中的公式、图片、形状；失败则返回 False。"""
    try:
        import xlwings as xw
    except Exception as error:
        print(f"[法兰报告][xlwings] {label}不可用，使用 openpyxl：{error}")
        return False

    app = None
    workbook = None
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        app = xw.App(visible=False, add_book=False)
        app.display_alerts = False
        app.screen_updating = False
        workbook = app.books.open(
            str(template_path),
            update_links=False,
            read_only=False,
        )
        fill_workbook(workbook)
        workbook.save(str(output_path))
        print(f"[法兰报告][xlwings] {label}已生成：{output_path}")
        return True
    except Exception as error:
        print(f"[法兰报告][xlwings] {label}失败，回退 openpyxl：{error}")
        traceback.print_exc()
        return False
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
        if app is not None:
            try:
                app.quit()
            except Exception:
                pass


def _generate_parameter_overview_with_xlwings(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> bool:
    datas = _get_flange_datas(json_data)

    def _fill(workbook):
        sheet = _xlwings_sheet(workbook, FLANGE_MODULE_NAME)
        last_row = max(2, sheet.used_range.last_cell.row)
        sheet.range((2, 1), (last_row, 3)).clear_contents()
        rows = [
            [
                item.get("Id", ""),
                item.get("Name", ""),
                item.get("Value", ""),
            ]
            for item in datas
        ]
        if rows:
            sheet.range((2, 1)).value = rows

    return _try_generate_with_xlwings(
        template_path,
        output_path,
        _fill,
        "参数总览表",
    )


def _generate_parameter_overview_with_openpyxl(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> None:
    """openpyxl 兜底：保留单元格样式，但部分图片/形状可能无法保留。"""
    datas = _get_flange_datas(json_data)
    workbook = load_workbook(template_path)
    try:
        sheet = workbook[FLANGE_MODULE_NAME] if FLANGE_MODULE_NAME in workbook.sheetnames else workbook.active
        for row in sheet.iter_rows(min_row=2, max_col=3):
            for cell in row:
                cell.value = None
        for row_index, item in enumerate(datas, start=2):
            sheet.cell(row=row_index, column=1, value=item.get("Id", ""))
            sheet.cell(row=row_index, column=2, value=item.get("Name", ""))
            sheet.cell(row=row_index, column=3, value=item.get("Value", ""))
        workbook.save(output_path)
    finally:
        workbook.close()


def generate_parameter_overview(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> None:
    """优先用 xlwings 生成参数总览表，保留公式和图像；不可用时回退 openpyxl。"""
    if _generate_parameter_overview_with_xlwings(json_data, template_path, output_path):
        return
    _generate_parameter_overview_with_openpyxl(json_data, template_path, output_path)


def _report_labels(sheet, row_index: int) -> Tuple[str, ...]:
    return tuple(
        label
        for label in (
            _stringify_value(sheet.cell(row=row_index, column=1).value),
            _stringify_value(sheet.cell(row=row_index, column=3).value),
        )
        if _normalise_key(label) not in GENERIC_REPORT_LABELS
    )


def _text_score(labels: Sequence[str], source_item: Dict[str, Any]) -> int:
    """以模板行的文本和 JSON Name/Desc 计算匹配分，不依赖硬编码字段表。"""
    source_labels = (
        _stringify_value(source_item.get("Name")),
        _stringify_value(source_item.get("Desc")),
    )
    score = 0
    for report_label in labels:
        report_key = _normalise_key(report_label)
        if not report_key:
            continue
        for source_label in source_labels:
            source_key = _normalise_key(source_label)
            if not source_key:
                continue
            if report_key == source_key:
                score = max(score, 1000)
            elif len(report_key) >= 3 and (report_key in source_key or source_key in report_key):
                score = max(score, 500 + min(len(report_key), len(source_key)))
            else:
                report_pairs = {report_key[i:i + 2] for i in range(len(report_key) - 1)}
                source_pairs = {source_key[i:i + 2] for i in range(len(source_key) - 1)}
                score = max(score, len(report_pairs & source_pairs) * 20)
    return score


def _previous_specific_labels(sheet, row_index: int) -> Tuple[str, ...]:
    """为“允许值/校核结果”这类泛化字段取最近的上文参数名称。"""
    for previous_row in range(row_index - 1, 4, -1):
        row_name = _stringify_value(sheet.cell(row=previous_row, column=1).value)
        row_field = _stringify_value(sheet.cell(row=previous_row, column=3).value)
        if (
            "允许值" in row_field
            or "许用值" in row_field
            or "校核结果" in row_field
            or "待补充" in row_name
        ):
            continue
        labels = _report_labels(sheet, previous_row)
        if labels:
            return labels
    return ()


def _dynamic_report_value(sheet, row_index: int, datas: List[Dict[str, Any]]) -> Optional[str]:
    """通过报告模板字段与 JSON 的 Name/Desc 动态关联，不维护字段映射常量。"""
    labels = _report_labels(sheet, row_index)
    raw_labels = (
        _stringify_value(sheet.cell(row=row_index, column=1).value),
        _stringify_value(sheet.cell(row=row_index, column=3).value),
    )
    is_allowance = any("允许值" in label or "许用值" in label for label in raw_labels)
    is_check_result = any("校核结果" in label for label in raw_labels)
    context_labels = _previous_specific_labels(sheet, row_index) if is_check_result else ()

    scored: List[Tuple[int, Dict[str, Any]]] = []
    for item in datas:
        source_name = _stringify_value(item.get("Name"))
        if not source_name:
            continue
        score = _text_score(labels, item)
        source_key = _normalise_key(source_name)

        # 模板中的“允许值”不携带材料名称，按 JSON 内的语义自动选取许用应力。
        if is_allowance and "许用应力" in source_name:
            score += 180
            if "法兰" in source_name:
                score += 80
            if "设计温度" in source_name:
                score += 40

        # “校核结果”使用前一行的应力/刚度等上下文自动匹配 JSON 的同类校核项。
        if is_check_result and ("校核" in source_name or "结果" in source_name):
            context_score = _text_score(context_labels, item)
            score += context_score + 120
        if source_key and score:
            scored.append((score, item))

    if not scored:
        return None
    scored.sort(key=lambda pair: pair[0], reverse=True)
    best_score, best_item = scored[0]
    if best_score < 60:
        return None
    # 同分且值不同代表语义不够明确，宁可留空也不写错值。
    if len(scored) > 1 and scored[1][0] == best_score:
        if _stringify_value(scored[1][1].get("Value")) != _stringify_value(best_item.get("Value")):
            return None
    return _stringify_value(best_item.get("Value"))


def _generate_calculation_report_with_xlwings(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> bool:
    datas = _get_flange_datas(json_data)
    value_map = _parameter_value_map(datas)
    module_data = _get_flange_module_data(json_data)

    def _fill(workbook):
        sheet = _xlwings_sheet(workbook, FLANGE_MODULE_NAME)
        last_row = max(5, sheet.used_range.last_cell.row)
        names = sheet.range((5, 1), (last_row, 1)).value
        if not isinstance(names, list):
            names = [names]

        for offset, raw_name in enumerate(names, start=5):
            report_parameter_name = _stringify_value(raw_name)
            if not report_parameter_name:
                continue
            value = _find_value(value_map, (report_parameter_name,))
            if value is not None:
                sheet.range((offset, 4)).value = value

        sheet.range((126, 4)).value = "合格" if module_data.get("IsSuccess") else "不合格"

    return _try_generate_with_xlwings(
        template_path,
        output_path,
        _fill,
        "计算报告",
    )


def _generate_calculation_report_with_openpyxl(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> None:
    """openpyxl 兜底：严格按“JSON 参数名 == 计算报告 A 列”填充计算报告。"""
    datas = _get_flange_datas(json_data)
    value_map = _parameter_value_map(datas)
    workbook = load_workbook(template_path)
    try:
        sheet = workbook[FLANGE_MODULE_NAME] if FLANGE_MODULE_NAME in workbook.sheetnames else workbook.active
        for row_index in range(5, sheet.max_row + 1):
            report_parameter_name = _stringify_value(sheet.cell(row=row_index, column=1).value)
            if not report_parameter_name:
                continue
            value = _find_value(value_map, (report_parameter_name,))
            if value is not None:
                sheet.cell(row=row_index, column=4, value=value)

        module_data = _get_flange_module_data(json_data)
        sheet.cell(row=126, column=4, value="合格" if module_data.get("IsSuccess") else "不合格")
        workbook.save(output_path)
    finally:
        workbook.close()


def generate_calculation_report(
    json_data: Dict[str, Any], template_path: Path, output_path: Path
) -> None:
    """优先用 xlwings 生成计算报告，保留公式和图像；不可用时回退 openpyxl。"""
    if _generate_calculation_report_with_xlwings(json_data, template_path, output_path):
        return
    _generate_calculation_report_with_openpyxl(json_data, template_path, output_path)


def _coating_required(value_map: Dict[str, str]) -> bool:
    thickness = _find_value(value_map, ("覆层厚度", "法兰覆层厚度"))
    if thickness is None:
        return False
    text = thickness.strip()
    if text in {"", "0", "0.0", "0.00", "否", "none"}:
        return False
    try:
        return float(text) != 0
    except ValueError:
        return True


def _close_open_target_drawing(acad, target_path: Path) -> None:
    target = os.path.normcase(os.path.normpath(str(target_path.resolve())))
    documents = []
    try:
        for index in range(acad.Documents.Count):
            document = acad.Documents.Item(index)
            full_name = getattr(document, "FullName", "") or ""
            if full_name and os.path.normcase(os.path.normpath(full_name)) == target:
                documents.append(document)
    except Exception:
        return
    for document in documents:
        try:
            document.Close(False)
        except Exception:
            pass


def _write_cad_value(
    doc, value_map: Dict[str, str], handle: str, *names: str, decimal_places: Optional[int] = None
) -> None:
    value = _find_value(value_map, names)
    if value is None:
        return
    if decimal_places is not None:
        try:
            value = f"{float(value):.{decimal_places}f}"
        except (TypeError, ValueError):
            pass
    from modules.TwoD.TwoD_peizhi import safe_modify

    safe_modify(doc, handle, value)


def generate_flange_drawing(
    value_map: Dict[str, str], save_dir: Path, root: Path, timestamp_label: str
) -> Path:
    """调用 TwoD 的 CAD 接口，所有尺寸均由单独法兰 JSON 提供。"""
    from modules.TwoD.TwoD_peizhi import (
        get_autocad_instance,
        open_drawing_with_wait,
        refresh_doc,
        safe_modify,
        wait_cad_idle,
    )
    from modules.TwoD.cad_context import CAD_CTX

    coated = _coating_required(value_map)
    template_path = root / ("法兰凹-覆层.dwg" if coated else "法兰-凹.dwg")
    if not template_path.is_file():
        raise FileNotFoundError(f"未找到 CAD 模板：{template_path}")

    acad = get_autocad_instance()
    if acad is None:
        raise RuntimeError("未检测到可用 AutoCAD，无法生成法兰图纸。")

    save_dir.mkdir(parents=True, exist_ok=True)
    target_path = save_dir / f"法兰_{timestamp_label}.dwg"
    _close_open_target_drawing(acad, target_path)
    shutil.copy2(template_path, target_path)
    wait_cad_idle(acad, timeout=20)
    acad, document = open_drawing_with_wait(str(target_path))
    if document is None:
        raise RuntimeError(f"无法打开 CAD 图纸：{target_path}")

    CAD_CTX.update({"product_id": None, "save_dir": str(save_dir), "target_dwg": str(target_path), "doc": document})
    wait_cad_idle(acad, timeout=30)
    refresh_doc(document)

    if coated:
        handle_map = {
            "325f": ("法兰名义外径",), "3260": ("法兰名义内径",), "3282": ("D2",),
            "3289": ("D3",), "3263": ("法兰名义厚度",), "3264": ("法兰颈部高度",),
            "3265": ("法兰总高",), "3267": ("螺栓数量",), "3268": ("螺栓孔直径",),
            "326d": ("法兰直边段高度",), "3261": ("螺栓中心圆直径",),
            "332b": ("法兰材料牌号", "材料牌号"), "18ed": ("覆层厚度", "法兰覆层厚度"),
            "18db": ("R", "圆角半径", "法兰圆角半径"), "332c": ("法兰毛坯质量",),
            "3385": ("法兰成型质量",),
        }
        small_handle, large_handle, title_handles = "3262", "326a", ("326e", "1a25")
    else:
        handle_map = {
            "15AE": ("法兰名义外径",), "15AD": ("法兰名义内径",), "130f3": ("D2",),
            "130fa": ("D3",), "15B4": ("法兰名义厚度",), "15B5": ("法兰颈部高度",),
            "15B6": ("法兰总高",), "15B8": ("螺栓数量",), "15B9": ("螺栓孔直径",),
            "15C9": ("法兰直边段高度",), "15AF": ("螺栓中心圆直径",),
            "1319c": ("法兰材料牌号", "材料牌号"), "15B7": ("R", "圆角半径", "法兰圆角半径"),
            "1319d": ("法兰毛坯质量",), "131fe": ("法兰成型质量",),
        }
        small_handle, large_handle, title_handles = "15B2", "15B3", ("18CB", "1B11")

    for handle, names in handle_map.items():
        _write_cad_value(
            document,
            value_map,
            handle,
            *names,
            decimal_places=2 if any("质量" in name for name in names) else None,
        )
    inner_diameter = _find_value(value_map, ("法兰名义内径",))
    small_thickness = _find_value(value_map, ("法兰颈部小端名义厚度",))
    large_thickness = _find_value(value_map, ("法兰颈部大端名义厚度",))
    try:
        safe_modify(document, small_handle, float(inner_diameter) + 2 * float(small_thickness))
        safe_modify(document, large_handle, float(inner_diameter) + 2 * float(large_thickness))
    except (TypeError, ValueError):
        pass
    for handle in title_handles:
        safe_modify(document, handle, "法兰")
    try:
        document.Regen(1)
    except Exception:
        pass
    document.Save()
    return target_path


def run_standalone_flange_report(parent=None, program_root: Optional[str] = None) -> None:
    """菜单入口：选择输出目录后，基于法兰独立计算 Output JSON 生成报告和图纸。"""
    root = _program_root(program_root)
    print(f"[法兰流程] 程序根目录：{root}")
    save_dir_text = QFileDialog.getExistingDirectory(parent, "选择法兰报告/图纸保存目录", str(root))
    if not save_dir_text:
        return

    try:
        save_dir = Path(save_dir_text).resolve()
        output_json_path, json_data = load_flange_output_json(root)
        timestamp_label = datetime.fromtimestamp(
            output_json_path.stat().st_mtime
        ).strftime("%Y-%m-%d-%H-%M-%S")

        overview_template = root / INTERMEDIATE_TEMPLATE_FILENAME
        report_template = root / REPORT_TEMPLATE_FILENAME
        if not overview_template.is_file() or not report_template.is_file():
            raise FileNotFoundError("未找到法兰参数总览表或计算报告模板。")

        overview_path = save_dir / f"强度计算元件输出参数表_法兰_{timestamp_label}.xlsx"
        report_path = save_dir / f"计算报告（法兰）_{timestamp_label}.xlsx"
        generate_parameter_overview(json_data, overview_template, overview_path)
        generate_calculation_report(json_data, report_template, report_path)

        drawing_path = None
        drawing_error = None
        try:
            drawing_path = generate_flange_drawing(
                _parameter_value_map(_get_flange_datas(json_data)), save_dir, root, timestamp_label
            )
        except Exception as error:
            drawing_error = error
            traceback.print_exc()

        message = (
            f"已使用单独法兰计算输出：\n{output_json_path}\n\n"
            f"参数总览表：\n{overview_path}\n\n计算报告：\n{report_path}"
        )
        if drawing_path:
            message += f"\n\n图纸：\n{drawing_path}"
        if drawing_error:
            QMessageBox.warning(parent, "法兰单独报告/图纸", message + f"\n\n图纸生成失败：\n{drawing_error}")
        else:
            QMessageBox.information(parent, "法兰单独报告/图纸", message)
    except Exception as error:
        traceback.print_exc()
        QMessageBox.critical(parent, "法兰单独报告/图纸", f"生成失败：\n{error}")
